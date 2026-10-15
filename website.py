import streamlit as st
import pandas as pd
import io
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="ระบบประมวลผลและวิเคราะห์ราคา", layout="wide", initial_sidebar_state="collapsed")

st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Kanit:wght@300;400;500;600&display=swap');

    .banner-main-title, .banner-sub-title, .card-header, .card-subheader, 
    .sub-card-green, .sub-card-red, .stat-number, .stat-label, .stat-unit, h1, h2, h3 {
        font-family: 'Kanit', sans-serif !important;
    }

    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}

    .banner-container {
        background-color: #0d2137;
        border-radius: 16px;
        padding: 35px 40px;
        margin-bottom: 24px;
        box-shadow: 0 4px 12px rgba(0,0,0,0.08);
        color: white;
    }
    .banner-main-title {
        font-size: 26px;
        font-weight: 600;
        color: #ffffff;
        margin-bottom: 4px;
    }
    .banner-sub-title {
        font-size: 15px;
        color: #9fb3c8;
    }

    .dashboard-card {
        background-color: #f0f4f8;
        border: 1px solid #d9e2ec;
        border-radius: 16px;
        padding: 24px;
        box-shadow: 0 4px 12px rgba(0, 0, 0, 0.03);
        margin-bottom: 20px;
    }
    .card-header {
        font-size: 18px;
        font-weight: 600;
        color: #102a43;
        margin-bottom: 4px;
    }
    .card-subheader {
        font-size: 13px;
        color: #627d98;
        margin-bottom: 16px;
    }

    .sub-card-green {
        background-color: #c5e1a5;
        border-radius: 12px;
        padding: 20px;
        text-align: center;
        flex: 1;
        color: #33691e;
    }

    .sub-card-red {
        background-color: #f9bdbb;
        border-radius: 12px;
        padding: 20px;
        text-align: center;
        flex: 1;
        color: #b0120a;
    }

    .stat-number {
        font-size: 38px;
        font-weight: 600;
        line-height: 1.2;
        margin: 8px 0;
    }
    .stat-label {
        font-size: 15px;
        font-weight: 600;
    }
    .stat-unit {
        font-size: 13px;
        opacity: 0.85;
    }
    </style>
""", unsafe_allow_html=True)

st.markdown("""
    <div class="banner-container">
        <div class="banner-main-title">📊 ระบบประมวลผลและวิเคราะห์ราคาสินค้าสำคัญ</div>
        <div class="banner-sub-title">อัปโหลดไฟล์ Excel ผลลัพธ์ของคุณเพื่อประมวลผล</div>
    </div>
""", unsafe_allow_html=True)

uploaded_file = st.file_uploader("📂 เลือกไฟล์ Excel ผลลัพธ์ (.xlsx)", type=["xlsx"])

if uploaded_file is not None:
    try:
        df = pd.read_excel(uploaded_file)
        
        # ค้นหาคอลัมน์ปี พ.ศ. 4 หลักแบบอัตโนมัติ
        year_columns = []
        for col in df.columns:
            col_str = str(col).strip()
            if re.match(r'^\d{4}$', col_str):
                year_columns.append(col)
                
        year_columns = sorted(year_columns, key=lambda x: int(str(x)))
        
        if len(year_columns) < 2:
            st.error("⚠️ ไม่พบคอลัมน์ปี พ.ศ. อย่างน้อย 2 ปีในไฟล์ Excel กรุณาตรวจสอบหัวตาราง")
            st.stop()
            
        prev_year = year_columns[-2]
        current_year = year_columns[-1]
        
        st.success(f"✅ อัปโหลดและอ่านไฟล์สำเร็จ! ระบบกำลังเปรียบเทียบข้อมูลปี {prev_year} และ {current_year}")
        
        def assign_unit(name):
            name_str = str(name).strip()
            if "ไข่ไก่" in name_str:
                return "บาท/ฟอง"
            elif "มะนาว" in name_str:
                return "บาท/ผล"
            else:
                return "บาท/กิโลกรัม"
                
        df['หน่วย'] = df['COMMODITY_NAME'].apply(assign_unit)
        
        name_mapping = {
            "ไข่ไก่": "ไข่ไก่ (เบอร์ 2)",
            "เนื้อสุกร": "เนื้อสุกร (สามชั้นตัดตามยาว)",
            "ไก่สด": "ไก่สด (ทั้งตัวไม่รวมเครื่องใน)"
        }
        df['COMMODITY_NAME'] = df['COMMODITY_NAME'].replace(name_mapping)
        
        def fix_previous_year_price(row):
            p_curr = row.get(current_year, 0)
            p_prev = row.get(prev_year, 0)
            yoy = row.get('YoY', 0)
            
            if pd.notna(p_curr) and pd.notna(p_prev) and pd.notna(yoy):
                is_yoy_positive = yoy > 0
                is_price_increasing = p_curr > p_prev
                
                if is_yoy_positive != is_price_increasing and (1 + yoy / 100) != 0:
                    return round(p_curr / (1 + yoy / 100), 2)
            return p_prev

        # --- จัดการหน่วยและคำนวณราคาไข่ไก่ (รองรับทั้งแบบแผงและแบบฟอง) ---
        egg_mask = df['COMMODITY_NAME'] == "ไข่ไก่ (เบอร์ 2)"
        
        # 1. เช็กและแปลงปีปัจจุบัน (current_year) ให้เป็นราคาต่อฟอง (ถ้ามากกว่า 20 แสดงว่าเป็นราคาแผง)
        if df.loc[egg_mask, current_year].values.size > 0:
            curr_val = df.loc[egg_mask, current_year].values[0]
            if pd.notna(curr_val) and curr_val > 20:
                df.loc[egg_mask, current_year] = (df.loc[egg_mask, current_year] / 30).round(2)

        # 2. คำนวณราคาย้อนกลับสำหรับปีปีก่อน (prev_year) จากสูตร YoY
        df[prev_year] = df.apply(fix_previous_year_price, axis=1)

        # 3. เช็กและแปลงปีปีก่อน (prev_year) ให้เป็นราคาต่อฟองเช่นกัน (ถ้ามากกว่า 20 แสดงว่าเป็นราคาแผง)
        if df.loc[egg_mask, prev_year].values.size > 0:
            prev_val = df.loc[egg_mask, prev_year].values[0]
            if pd.notna(prev_val) and prev_val > 20:
                df.loc[egg_mask, prev_year] = (df.loc[egg_mask, prev_year] / 30).round(2)
        
        df['Group'] = df.apply(lambda r: 'เขียว' if (r['YoY'] >= 0 and r.get('ContributeYoY', 0) >= 0) else 'แดง', axis=1)
        
        priority_items = [
            "เนื้อสุกร (สามชั้นตัดตามยาว)", 
            "ไก่สด (ทั้งตัวไม่รวมเครื่องใน)", 
            "ไข่ไก่ (เบอร์ 2)"
        ]
        
        contribute_col = 'ContributeYoY' if 'ContributeYoY' in df.columns else (df.columns[df.columns.str.contains('Contribute', case=False)].tolist() or ['YoY'])[0]
        
        df['IsPriority'] = df['COMMODITY_NAME'].apply(lambda x: 0 if x in priority_items else 1)
        
        df_green = df[df['Group'] == 'เขียว'].copy()
        df_red = df[df['Group'] == 'แดง'].copy()
        
        df_green = df_green.sort_values(by=['IsPriority', contribute_col], ascending=[True, False])
        df_red = df_red.sort_values(by=['IsPriority', contribute_col], ascending=[True, True])
        
        df = pd.concat([df_green, df_red]).reset_index(drop=True)
        
        df_for_count = df[~df['COMMODITY_NAME'].isin(priority_items)]
        
        yoy_positive_count = len(df_for_count[df_for_count['YoY'] > 0])
        yoy_negative_count = len(df_for_count[df_for_count['YoY'] < 0])
        
        mom_positive_count = len(df_for_count[df_for_count['MoM'] > 0])
        mom_negative_count = len(df_for_count[df_for_count['MoM'] < 0])
        
        st.markdown("---")
        
        col1, col2 = st.columns(2)
        
        with col1:
            st.markdown(f"""
                <div class="dashboard-card">
                    <div class="card-header">📈 ราคาเทียบกับปีที่ผ่านมา (YoY)</div>
                    <div class="card-subheader"></div>
                    <div style="display: flex; gap: 15px;">
                        <div class="sub-card-green">
                            <div class="stat-label">🟢 ปรับราคาสูงขึ้น</div>
                            <div class="stat-number">{yoy_positive_count}</div>
                            <div class="stat-unit">รายการ</div>
                        </div>
                        <div class="sub-card-red">
                            <div class="stat-label">🔴 ปรับราคาลดลง</div>
                            <div class="stat-number">{yoy_negative_count}</div>
                            <div class="stat-unit">รายการ</div>
                        </div>
                    </div>
                </div>
            """, unsafe_allow_html=True)
            
        with col2:
            st.markdown(f"""
                <div class="dashboard-card">
                    <div class="card-header">📅 ราคาเทียบกับเดือนที่ผ่านมา (MoM)</div>
                    <div class="card-subheader"></div>
                    <div style="display: flex; gap: 15px;">
                        <div class="sub-card-green">
                            <div class="stat-label">🟢 ปรับราคาสูงขึ้น</div>
                            <div class="stat-number">{mom_positive_count}</div>
                            <div class="stat-unit">รายการ</div>
                        </div>
                        <div class="sub-card-red">
                            <div class="stat-label">🔴 ปรับราคาลดลง</div>
                            <div class="stat-number">{mom_negative_count}</div>
                            <div class="stat-unit">รายการ</div>
                        </div>
                    </div>
                </div>
            """, unsafe_allow_html=True)
            
        st.markdown("---")
        st.subheader(f"📋 ตารางแสดงรายการสินค้า (ปี {prev_year} - {current_year})")
        
        display_df = df[['COMMODITY_NAME', 'หน่วย', prev_year, current_year]].copy()
        display_df[prev_year] = display_df[prev_year].map('{:,.2f}'.format)
        display_df[current_year] = display_df[current_year].map('{:,.2f}'.format)
        
        def style_table(row):
            yoy_val = df.loc[row.name, 'YoY']
            item_name = str(df.loc[row.name, 'COMMODITY_NAME']).strip()
            
            bg_style = ''
            if pd.notna(yoy_val):
                if yoy_val > 0:
                    bg_style = 'background-color: #d4edda; font-family: Kanit, sans-serif;'
                elif yoy_val < 0:
                    bg_style = 'background-color: #f8d7da; font-family: Kanit, sans-serif;'
                else:
                    bg_style = 'font-family: Kanit, sans-serif;'
                    
            font_style = ''
            if item_name in priority_items:
                font_style = ' font-weight: bold;'
                
            combined = bg_style + font_style
            return [combined] * len(row)

        styled_table = display_df.style.apply(style_table, axis=1)
        st.dataframe(styled_table, use_container_width=True)
        
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            export_df = df[['COMMODITY_NAME', 'หน่วย', prev_year, current_year]].copy()
            export_df.columns = ['ชื่อสินค้า', 'หน่วย', str(prev_year), str(current_year)]
            export_df.to_excel(writer, index=False, sheet_name='Price_Report')
            
            workbook = writer.book
            worksheet = writer.sheets['Price_Report']
            
            green_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
            red_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )
            
            header_fill = PatternFill(start_color="102A43", end_color="102A43", fill_type="solid")
            header_font = Font(name="Kanit", size=11, bold=True, color="FFFFFF")
            for col_num in range(1, len(export_df.columns) + 1):
                cell = worksheet.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            for row_idx in range(len(export_df)):
                excel_row = row_idx + 2
                yoy_val = df.iloc[row_idx]['YoY']
                item_name = str(df.iloc[row_idx]['COMMODITY_NAME']).strip()
                
                row_fill = None
                if pd.notna(yoy_val):
                    if yoy_val > 0:
                        row_fill = green_fill
                    elif yoy_val < 0:
                        row_fill = red_fill
                        
                is_bold = item_name in priority_items
                row_font = Font(name="Kanit", size=11, bold=is_bold)
                
                for col_num in range(1, len(export_df.columns) + 1):
                    cell = worksheet.cell(row=excel_row, column=col_num)
                    if row_fill:
                        cell.fill = row_fill
                    cell.font = row_font
                    cell.border = thin_border
                    
                    if col_num in [3, 4]:
                        cell.number_format = '#,##0.00'
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    elif col_num == 2:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                        
            for col in worksheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                worksheet.column_dimensions[col_letter].width = max(max_len + 4, 15)
                
        excel_data = output.getvalue()
        
        st.download_button(
            label="📥 ดาวน์โหลดไฟล์ Excel รายงานผลลัพธ์",
            data=excel_data,
            file_name=f"price_report_{prev_year}_{current_year}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        
    except Exception as e:
        st.error(f"⚠️ เกิดข้อผิดพลาดในการประมวลผลไฟล์: {e}")